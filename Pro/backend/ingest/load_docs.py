from __future__ import annotations

import re
import unicodedata
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader


def _normalize(text: str) -> str:
    text = unicodedata.normalize("NFC", text)
    text = text.replace("\u00a0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def load_txt(path: Path) -> list[dict]:
    content = path.read_text(encoding="utf-8", errors="ignore")
    return [{"page": 1, "text": _normalize(content)}]


def load_pdf(path: Path) -> list[dict]:
    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        pages.append({"page": i, "text": _normalize(text)})
    return pages


def load_docx(path: Path) -> list[dict]:
    doc = Document(str(path))
    parts: list[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text.strip())
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                # Convert table row to "col: value" style pairs when possible
                if len(cells) == 2:
                    parts.append(f"{cells[0]}: {cells[1]}")
                else:
                    parts.append(" | ".join(cells))
    return [{"page": 1, "text": _normalize("\n".join(parts))}]


def load_html(path: Path) -> list[dict]:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    soup = BeautifulSoup(raw, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()
    # Tables → column: value lines
    for table in soup.find_all("table"):
        lines = []
        for row in table.find_all("tr"):
            cells = [c.get_text(" ", strip=True) for c in row.find_all(["th", "td"])]
            cells = [c for c in cells if c]
            if len(cells) >= 2:
                lines.append(f"{cells[0]}: {' | '.join(cells[1:])}")
            elif cells:
                lines.append(cells[0])
        table.replace_with("\n".join(lines) + "\n")
    text = soup.get_text("\n")
    return [{"page": 1, "text": _normalize(text)}]


LOADERS = {
    ".txt": load_txt,
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".html": load_html,
    ".htm": load_html,
}


def load_document(path: Path) -> list[dict]:
    suffix = path.suffix.lower()
    if suffix not in LOADERS:
        raise ValueError(f"Unsupported file type: {suffix}")
    return LOADERS[suffix](path)
